"""Word documents are read with their tables, in reading order."""

from __future__ import annotations

from pathlib import Path

from docx import Document

from bristlenose.utils.docx_text import block_texts


def _saved(tmp_path: Path, doc) -> Document:  # type: ignore[no-untyped-def]
    path = tmp_path / "d.docx"
    doc.save(str(path))
    return Document(str(path))


def test_table_text_is_read_in_place(tmp_path: Path) -> None:
    doc = Document()
    doc.add_paragraph("before")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "a1"
    table.cell(0, 1).text = "b1"
    table.cell(1, 0).text = "a2"
    table.cell(1, 1).text = "b2"
    doc.add_paragraph("after")

    assert list(block_texts(_saved(tmp_path, doc))) == [
        "before", "a1", "b1", "a2", "b2", "after",
    ]


def test_a_merged_cell_is_read_once(tmp_path: Path) -> None:
    doc = Document()
    table = doc.add_table(rows=1, cols=3)
    merged = table.cell(0, 0).merge(table.cell(0, 1))
    merged.text = "wide"
    table.cell(0, 2).text = "narrow"

    assert list(block_texts(_saved(tmp_path, doc))) == ["wide", "narrow"]


def test_a_nested_table_is_read(tmp_path: Path) -> None:
    doc = Document()
    outer = doc.add_table(rows=1, cols=1).cell(0, 0)
    outer.text = "outer"
    outer.add_table(rows=1, cols=1).cell(0, 0).text = "inner"

    assert list(block_texts(_saved(tmp_path, doc))) == ["outer", "inner", ""]
